import smtplib
import ssl
import time
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from openpyxl import load_workbook


def send_html_email(
    sender_email,
    sender_password,
    recipient_email,
    subject,
    html_content,
    plain_text=None,
    smtp_server="smtps.hiworks.com",
    smtp_port=465,
):
    # 이메일 메시지 구성
    message = MIMEMultipart("alternative")
    message["From"] = sender_email
    message["To"] = recipient_email
    message["Subject"] = subject

    # 일반 텍스트 버전 (HTML을 지원하지 않는 클라이언트용)
    if plain_text is None:
        plain_text = "이 이메일을 보려면 HTML을 지원하는 이메일 클라이언트가 필요합니다."

    part1 = MIMEText(plain_text, "plain")
    part2 = MIMEText(html_content, "html")

    # 순서 중요: 마지막에 추가된 것이 우선 표시됨
    message.attach(part1)
    message.attach(part2)

    # SMTP 서버 연결 및 전송 (Hiworks SSL - OpenSSL 3.0 호환)
    context = ssl.SSLContext(ssl.PROTOCOL_TLS_CLIENT)
    context.set_ciphers('DEFAULT:@SECLEVEL=1')
    context.check_hostname = False
    context.verify_mode = ssl.CERT_NONE

    with smtplib.SMTP_SSL(smtp_server, smtp_port, context=context) as server:
        server.login(sender_email, sender_password)
        server.sendmail(sender_email, recipient_email, message.as_string())

    print(f"[성공] {recipient_email}")


def send_bulk_from_excel(
    sender_email,
    sender_password,
    excel_path,
    subject,
    html_template,
    email_col="B",
    name_col="A",
    start_row=2,
    delay=1,
    smtp_server="smtps.hiworks.com",
    smtp_port=465,
):
    """
    엑셀 파일에서 수신자 목록을 읽어 이메일을 발송합니다.

    Args:
        sender_email: 발신자 이메일
        sender_password: 발신자 비밀번호
        excel_path: 엑셀 파일 경로
        subject: 이메일 제목 ({name} 치환 가능)
        html_template: HTML 템플릿 ({name}, {email} 치환 가능)
        email_col: 이메일 컬럼 (기본: B)
        name_col: 이름 컬럼 (기본: A)
        start_row: 시작 행 (기본: 2, 헤더 제외)
        delay: 발송 간격 초 (기본: 1초)
        smtp_server: SMTP 서버 (기본: smtps.hiworks.com)
        smtp_port: SMTP 포트 (기본: 465)

    Returns:
        {"success": 성공 건수, "fail": 실패 건수}
    """
    wb = load_workbook(excel_path)
    ws = wb.active

    success_count = 0
    fail_count = 0

    for row in range(start_row, ws.max_row + 1):
        name = ws[f"{name_col}{row}"].value
        email = ws[f"{email_col}{row}"].value

        if not email:
            continue

        # 템플릿 치환
        html_content = html_template.replace("{name}", str(name or ""))
        html_content = html_content.replace("{email}", str(email))
        email_subject = subject.replace("{name}", str(name or ""))

        try:
            send_html_email(
                sender_email=sender_email,
                sender_password=sender_password,
                recipient_email=email,
                subject=email_subject,
                html_content=html_content,
                smtp_server=smtp_server,
                smtp_port=smtp_port,
            )
            success_count += 1
            time.sleep(delay)  # 스팸 방지 딜레이
        except Exception as e:
            print(f"[실패] {email}: {e}")
            fail_count += 1

    print(f"\n완료: 성공 {success_count}건, 실패 {fail_count}건")
    wb.close()
    return {"success": success_count, "fail": fail_count}
