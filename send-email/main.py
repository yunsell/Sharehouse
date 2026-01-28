import os
import smtplib
import ssl
import time
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from openpyxl import load_workbook
from dotenv import load_dotenv

load_dotenv()


def send_html_email(sender_email, sender_password, recipient_email, subject, html_content, plain_text=None):
    # 이메일 메시지 구성
    message = MIMEMultipart("alternative")
    message["From"] = sender_email
    message["To"] = recipient_email
    message["Subject"] = subject

    # 일반 텍스트 버전 (HTML을 지원하지 않는 클라이언트용)
    if plain_text is None:
        plain_text = "이 이메일을 보려면 HTML을 지원하는 이메일 클라이언트가 필요합니다."
    
    part1 = MIMEText(plain_text, "plain")
    part2 = MIMEText(html_content, "html")

    # 순서 중요: 마지막에 추가된 것이 우선 표시됨
    message.attach(part1)
    message.attach(part2)

    # SMTP 서버 연결 및 전송 (Hiworks SSL - OpenSSL 3.0 호환)
    context = ssl.SSLContext(ssl.PROTOCOL_TLS_CLIENT)
    context.set_ciphers('DEFAULT:@SECLEVEL=1')
    context.check_hostname = False
    context.verify_mode = ssl.CERT_NONE

    smtp_server = os.getenv("SMTP_SERVER", "smtps.hiworks.com")
    smtp_port = int(os.getenv("SMTP_PORT", 465))

    with smtplib.SMTP_SSL(smtp_server, smtp_port, context=context) as server:
        server.login(sender_email, sender_password)
        server.sendmail(sender_email, recipient_email, message.as_string())
    
    print(f"[성공] {recipient_email}")


def send_bulk_from_excel(
    sender_email,
    sender_password,
    excel_path,
    subject,
    html_template,
    email_col="B",
    name_col="A",
    start_row=2,
    delay=1,
):
    """
    엑셀 파일에서 수신자 목록을 읽어 이메일을 발송합니다.

    Args:
        sender_email: 발신자 이메일
        sender_password: 발신자 비밀번호
        excel_path: 엑셀 파일 경로
        subject: 이메일 제목 ({name} 치환 가능)
        html_template: HTML 템플릿 ({name}, {email} 치환 가능)
        email_col: 이메일 컬럼 (기본: B)
        name_col: 이름 컬럼 (기본: A)
        start_row: 시작 행 (기본: 2, 헤더 제외)
        delay: 발송 간격 초 (기본: 1초)
    """
    wb = load_workbook(excel_path)
    ws = wb.active

    success_count = 0
    fail_count = 0

    for row in range(start_row, ws.max_row + 1):
        name = ws[f"{name_col}{row}"].value
        email = ws[f"{email_col}{row}"].value

        if not email:
            continue

        # 템플릿 치환
        html_content = html_template.replace("{name}", str(name or ""))
        html_content = html_content.replace("{email}", str(email))
        email_subject = subject.replace("{name}", str(name or ""))

        try:
            send_html_email(
                sender_email=sender_email,
                sender_password=sender_password,
                recipient_email=email,
                subject=email_subject,
                html_content=html_content,
            )
            success_count += 1
            time.sleep(delay)  # 스팸 방지 딜레이
        except Exception as e:
            print(f"[실패] {email}: {e}")
            fail_count += 1

    print(f"\n완료: 성공 {success_count}건, 실패 {fail_count}건")
    wb.close()


# HTML 템플릿 예시 ({name}으로 이름 치환)
html_template = """
<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <style>
        body {
            font-family: 'Apple SD Gothic Neo', 'Malgun Gothic', sans-serif;
            background-color: #f4f4f4;
            margin: 0;
            padding: 20px;
        }
        .container {
            max-width: 600px;
            margin: 0 auto;
            background-color: #ffffff;
            border-radius: 10px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
            overflow: hidden;
        }
        .header {
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
            padding: 30px;
            text-align: center;
        }
        .header h1 {
            margin: 0;
            font-size: 24px;
        }
        .content {
            padding: 30px;
            color: #333333;
            line-height: 1.6;
        }
        .button {
            display: inline-block;
            background-color: #667eea;
            color: white;
            padding: 12px 30px;
            text-decoration: none;
            border-radius: 5px;
            margin-top: 20px;
        }
        .footer {
            background-color: #f8f8f8;
            padding: 20px;
            text-align: center;
            font-size: 12px;
            color: #888888;
        }
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>🎉 환영합니다!</h1>
        </div>
        <div class="content">
            <p>안녕하세요, <strong>{name}</strong>님!</p>
            <p>회원가입을 축하드립니다. 저희 서비스를 이용해 주셔서 감사합니다.</p>
            <p>아래 버튼을 클릭하여 이메일 인증을 완료해주세요.</p>
            <center>
                <a href="https://example.com/verify" class="button">이메일 인증하기</a>
            </center>
        </div>
        <div class="footer">
            <p>본 메일은 발신 전용입니다.</p>
            <p>© 2025 Your Company. All rights reserved.</p>
        </div>
    </div>
</body>
</html>
"""

# 사용 예시
if __name__ == "__main__":
    # 환경변수에서 설정 로드
    sender_email = os.getenv("SENDER_EMAIL")
    sender_password = os.getenv("SENDER_PASSWORD")
    email_subject = os.getenv("EMAIL_SUBJECT", "뉴스레터")
    excel_path = os.getenv("EXCEL_PATH", "recipients.xlsx")
    html_path = os.getenv("HTML_PATH", "email.html")

    # HTML 템플릿 읽기
    with open(html_path, "r", encoding="utf-8") as f:
        html_from_file = f.read()

    send_bulk_from_excel(
        sender_email=sender_email,
        sender_password=sender_password,
        excel_path=excel_path,
        subject=email_subject,
        html_template=html_from_file,
        name_col="A",   # 이름 컬럼
        email_col="B",  # 이메일 컬럼
        start_row=2,    # 데이터 시작 행 (1행은 헤더)
        delay=1,        # 발송 간격 (초)
    )