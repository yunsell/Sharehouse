import re
import time
import requests
from openpyxl import load_workbook

MASS_SEND_URL = "https://apis.aligo.in/send_mass/"
REMAIN_URL = "https://apis.aligo.in/remain/"
MASS_LIMIT = 500  # send_mass 1회 최대 건수
SMS_MAX_BYTES = 90  # 초과하면 LMS


def normalize_phone(value):
    """엑셀 셀 값을 '01012345678' 형태로 정리합니다. 올바르지 않으면 None."""
    if value is None:
        return None
    if isinstance(value, float):
        value = int(value)
    digits = re.sub(r"\D", "", str(value))
    # 엑셀에서 숫자로 저장되어 앞자리 0이 빠진 경우 (1012345678)
    if len(digits) in (9, 10) and digits.startswith("1"):
        digits = "0" + digits
    if not re.fullmatch(r"01\d{8,9}", digits):
        return None
    return digits


def msg_bytes(text):
    # 알리고는 EUC-KR 기준으로 길이를 계산 (한글 2바이트)
    return len(text.encode("cp949", errors="replace"))


def get_remain(api_key, user_id):
    """
    발송 가능 건수를 조회합니다. API 키, 아이디, 발송 서버 IP 등록 확인용으로도 사용합니다.

    Returns:
        알리고 응답 (result_code가 1이면 성공, SMS_CNT / LMS_CNT / MMS_CNT 포함)
    """
    response = requests.post(REMAIN_URL, data={"key": api_key, "userid": user_id}, timeout=10)
    return response.json()


def read_recipients(excel_path, name_col="A", phone_col="B", start_row=2):
    """엑셀에서 (이름, 전화번호) 목록을 읽습니다. 잘못된 번호와 중복은 제외합니다."""
    wb = load_workbook(excel_path, read_only=True, data_only=True)
    ws = wb.active

    recipients = []
    seen = set()
    for row in range(start_row, ws.max_row + 1):
        name = ws[f"{name_col}{row}"].value
        raw_phone = ws[f"{phone_col}{row}"].value

        if raw_phone is None:
            continue

        phone = normalize_phone(raw_phone)
        if phone is None:
            print(f"[건너뜀] {row}행: 전화번호 형식 오류 ({raw_phone})")
            continue
        if phone in seen:
            print(f"[건너뜀] {row}행: 중복 번호 ({phone})")
            continue

        seen.add(phone)
        recipients.append((str(name or "").strip(), phone))

    wb.close()
    return recipients


def send_bulk_sms(api_key, user_id, sender, recipients, msg_template, title="", msg_type=None, test_mode=True, delay=1):
    """
    수신자별로 {name}, {phone}을 치환한 문자를 send_mass API로 발송합니다.

    Args:
        api_key: 알리고 API 키
        user_id: 알리고 아이디
        sender: 알리고에 등록된 발신번호
        recipients: [(이름, 전화번호), ...]
        msg_template: 문자 내용 ({name}, {phone} 치환 가능)
        title: 제목 (LMS/MMS에만 적용)
        msg_type: SMS / LMS (None이면 가장 긴 메시지 길이로 자동 결정)
        test_mode: True면 알리고 테스트모드 (실제 발송·차감 안 됨)
        delay: 500건 단위 요청 사이 간격 (초)

    Returns:
        {"success": 성공 건수, "fail": 실패 건수, "msg_type": 발송 타입, "msg_ids": [알리고 메시지 ID, ...]}
    """
    messages = [
        (phone, msg_template.replace("{name}", name).replace("{phone}", phone))
        for name, phone in recipients
    ]

    if msg_type is None:
        longest = max(msg_bytes(msg) for _, msg in messages)
        msg_type = "LMS" if longest > SMS_MAX_BYTES else "SMS"

    success_count = 0
    fail_count = 0
    msg_ids = []

    for start in range(0, len(messages), MASS_LIMIT):
        chunk = messages[start:start + MASS_LIMIT]

        data = {
            "key": api_key,
            "userid": user_id,
            "sender": sender,
            "cnt": len(chunk),
            "msg_type": msg_type,
            "testmode_yn": "Y" if test_mode else "N",
        }
        if msg_type != "SMS" and title:
            data["title"] = title
        for i, (phone, msg) in enumerate(chunk, start=1):
            data[f"rec_{i}"] = phone
            data[f"msg_{i}"] = msg

        label = f"{start + 1}~{start + len(chunk)}번째"
        try:
            response = requests.post(MASS_SEND_URL, data=data, timeout=30)
            result = response.json()
        except Exception as e:
            print(f"[실패] {label}: {e}")
            fail_count += len(chunk)
            continue

        if str(result.get("result_code")) == "1":
            ok = int(result.get("success_cnt", len(chunk)))
            ng = int(result.get("error_cnt", 0))
            success_count += ok
            fail_count += ng
            msg_ids.append(result.get("msg_id"))
            print(f"[성공] {label}: 성공 {ok}건, 실패 {ng}건 (mid: {result.get('msg_id')})")
        else:
            fail_count += len(chunk)
            print(f"[실패] {label}: {result.get('message')} (result_code: {result.get('result_code')})")

        if start + MASS_LIMIT < len(messages):
            time.sleep(delay)

    print(f"\n완료: 성공 {success_count}건, 실패 {fail_count}건 ({msg_type})")
    return {"success": success_count, "fail": fail_count, "msg_type": msg_type, "msg_ids": msg_ids}
